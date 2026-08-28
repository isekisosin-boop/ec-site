from flask import Flask, render_template, request, redirect, url_for, session
import os
import resend
from dotenv import load_dotenv
load_dotenv()

resend.api_key = os.getenv("RESEND_API_KEY")
from datetime import datetime
from openpyxl import load_workbook

app = Flask(__name__, static_folder="static", static_url_path="/static")

# セッションを利用するための秘密鍵
app.secret_key = "ec-site-development-key"


# =========================
# Excel商品データ
# =========================

def load_products_from_excel():

    workbook = load_workbook("商品一覧.xlsx", data_only=True)
    sheet = workbook.active

    products = {}

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if not row[0]:
            continue

        product_id = int(row[0])

        products[product_id] = {
            "name": row[1],
            "category": row[2],
            "description": row[3],
            "price": row[4],
            "detail": row[5],
            "features": row[6],
            "usage": row[7],
            "notes": row[8],
            "image": f"images/{row[9]}.jpg" if row[9] else "",
            "seo_title": row[10],
            "seo_description": row[11],
            "status": row[12]
        }

    return {
        product_id: product
        for product_id, product in products.items()
        if str(product["status"]).strip() == "公開"
    }


products_from_excel = load_products_from_excel()

# =========================
# カテゴリー対応
# =========================

category_map = {
    "oral": "口腔ケア",
    "skin": "スキンケア",
    "shampoo": "シャンプー",
    "other": "その他"
}


# =========================
# ホーム
# =========================

@app.route("/")
def index():
    return render_template(
        "index.html",
        products=products_from_excel
    )


# =========================
# 商品一覧
# =========================

@app.route("/products")
def products():

    category = request.args.get("category")
    search_query = request.args.get("q", "").strip()
    sort_order = request.args.get("sort", "default")

    filtered_products = products_from_excel

    # -------------------------
    # カテゴリー絞り込み
    # -------------------------

    if category:
        category_name = category_map.get(category)

        if category_name:
            filtered_products = {
                product_id: product
                for product_id, product in filtered_products.items()
                if product["category"] == category_name
            }

    # -------------------------
    # 商品検索
    # -------------------------

    if search_query:
        search_text = search_query.lower()

        filtered_products = {
            product_id: product
            for product_id, product in filtered_products.items()
            if search_text in product["name"].lower()
            or search_text in product["category"].lower()
            or search_text in product["description"].lower()
            or search_text in product["detail"].lower()
        }

    # -------------------------
    # 並び順
    # -------------------------

    if sort_order == "name":
        filtered_products = dict(
            sorted(
                filtered_products.items(),
                key=lambda item: item[1]["name"]
            )
        )

    return render_template(
        "products.html",
        products=filtered_products,
        selected_category=category,
        search_query=search_query,
        sort_order=sort_order
    )


# =========================
# 商品詳細
# =========================

@app.route("/products/<int:product_id>")
def product_detail(product_id):

    product = products_from_excel.get(product_id)

    if product is None:
        return "商品が見つかりません", 404

    return render_template(
        "product_detail.html",
        product=product,
        product_id=product_id
    )


# =========================
# カートに追加
# =========================

@app.route("/cart/add/<int:product_id>", methods=["POST"])
def add_to_cart(product_id):

    product = products_from_excel.get(product_id)

    if product is None:
        return "商品が見つかりません", 404

    quantity = request.form.get("quantity", 1, type=int)

    if quantity < 1:
        quantity = 1

    cart = session.get("cart", {})

    product_key = str(product_id)

    if product_key in cart:
        cart[product_key] += quantity
    else:
        cart[product_key] = quantity

    session["cart"] = cart

    return redirect(url_for("cart"))


# =========================
# カート数量変更
# =========================

@app.route("/cart/update/<int:product_id>", methods=["POST"])
def update_cart(product_id):

    if product_id not in products_from_excel:
        return "商品が見つかりません", 404

    quantity = request.form.get("quantity", 1, type=int)

    cart = session.get("cart", {})
    product_key = str(product_id)

    if quantity <= 0:
        cart.pop(product_key, None)
    else:
        cart[product_key] = quantity

    session["cart"] = cart

    return redirect(url_for("cart"))


# =========================
# カートから商品削除
# =========================

@app.route("/cart/remove/<int:product_id>", methods=["POST"])
def remove_from_cart(product_id):

    cart = session.get("cart", {})
    product_key = str(product_id)

    cart.pop(product_key, None)

    session["cart"] = cart

    return redirect(url_for("cart"))


# =========================
# カート
# =========================

@app.route("/cart")
def cart():

    cart_data = session.get("cart", {})

    cart_items = []
    total_quantity = 0

    for product_id, quantity in cart_data.items():

        product = products_from_excel.get(int(product_id))

        if product is None:
            continue

        cart_items.append({
            "id": int(product_id),
            "product": product,
            "quantity": quantity
        })

        total_quantity += quantity

    return render_template(
        "cart.html",
        cart_items=cart_items,
        total_quantity=total_quantity
    )


# =========================
# 購入手続き
# =========================

@app.route("/checkout")
def checkout():

    cart_data = session.get("cart", {})

    if not cart_data:
        return redirect(url_for("cart"))

    return render_template(
        "checkout.html"
    )


# =========================
# 注文内容確認
# =========================

@app.route("/checkout/confirm", methods=["POST"])
def checkout_confirm():

    cart_data = session.get("cart", {})

    if not cart_data:
        return redirect(url_for("cart"))

    customer = {
        "name": request.form.get("name", "").strip(),
        "email": request.form.get("email", "").strip(),
        "postal_code": request.form.get("postal_code", "").strip(),
        "address": request.form.get("address", "").strip(),
        "phone": request.form.get("phone", "").strip()
    }

    # 入力内容をセッションに一時保存
    session["customer"] = customer

    cart_items = []
    total_quantity = 0

    for product_id, quantity in cart_data.items():

        product = products_from_excel.get(int(product_id))

        if product is None:
            continue

        cart_items.append({
            "id": int(product_id),
            "product": product,
            "quantity": quantity
        })

        total_quantity += quantity

    return render_template(
        "checkout_confirm.html",
        customer=customer,
        cart_items=cart_items,
        total_quantity=total_quantity
    )


# =========================
# 注文確定
# =========================

@app.route("/order/complete", methods=["POST"])
def order_complete():

    cart_data = session.get("cart", {})
    customer = session.get("customer")

    # カートまたはお客様情報がない場合
    if not cart_data or not customer:
        return redirect(url_for("cart"))

    cart_items = []
    total_quantity = 0

    for product_id, quantity in cart_data.items():

        product = products_from_excel.get(int(product_id))

        if product is None:
            continue

        cart_items.append({
            "id": int(product_id),
            "product": product,
            "quantity": quantity
        })

        total_quantity += quantity

    # 仮の注文番号を作成
    order_number = datetime.now().strftime("EC%Y%m%d%H%M%S")

    # 注文完了ページで使用する情報を一時保存
    session["completed_order"] = {
        "order_number": order_number,
        "customer": customer,
        "cart_items": [
            {
                "id": item["id"],
                "quantity": item["quantity"]
            }
            for item in cart_items
        ],
        "total_quantity": total_quantity
    }

    # 注文確定後はカートを空にする
    session.pop("cart", None)
    session.pop("customer", None)

    return redirect(url_for("order_complete_page"))


# =========================
# 注文完了ページ
# =========================

@app.route("/order/complete")
def order_complete_page():

    completed_order = session.get("completed_order")

    if not completed_order:
        return redirect(url_for("products"))

    cart_items = []

    for item in completed_order["cart_items"]:

        product = products_from_excel.get(item["id"])

        if product is None:
            continue

        cart_items.append({
            "id": item["id"],
            "product": product,
            "quantity": item["quantity"]
        })

    return render_template(
        "order_complete.html",
        order_number=completed_order["order_number"],
        customer=completed_order["customer"],
        cart_items=cart_items,
        total_quantity=completed_order["total_quantity"]
    )


# =========================
# ブランドについて
# =========================

@app.route("/about")
def about():
    return render_template("about.html")


# =========================
# よくある質問
# =========================

@app.route("/faq")
def faq():
    return render_template("faq.html")

# =========================
# お問い合わせ
# =========================

@app.route("/contact", methods=["GET", "POST"])
def contact():

    if request.method == "POST":

        name = request.form.get("name", "").strip()
        email = request.form.get("email", "").strip()
        phone = request.form.get("phone", "").strip()
        category = request.form.get("category", "").strip()
        message = request.form.get("message", "").strip()

        # 必須項目チェック
        if not name or not email or not category or not message:
            return render_template(
                "contact.html",
                error="必須項目を入力してください。",
                form_data=request.form
            )

        # -------------------------
        # メール送信
        # -------------------------

        email_body = f"""
        <h2>お問い合わせがありました</h2>

        <p><strong>お名前：</strong>{name}</p>

        <p><strong>メールアドレス：</strong>{email}</p>

        <p><strong>電話番号：</strong>{phone or "未入力"}</p>

        <p><strong>お問い合わせ項目：</strong>{category}</p>

        <p><strong>お問い合わせ内容：</strong></p>

        <p>{message.replace(chr(10), "<br>")}</p>
        """

        resend.Emails.send({
            "from": "onboarding@resend.dev",
            "to": ["isekisosin@gmail.com"],
            "subject": f"【EC SITE】お問い合わせ：{category}",
            "html": email_body,
        })

        return render_template(
            "contact_complete.html",
            name=name,
            email=email,
            category=category,
            message=message
        )

    return render_template("contact.html")
# =========================
# 会社概要・各種ページ
# =========================

@app.route("/company")
def company():
    return render_template("company.html")


@app.route("/terms")
def terms():
    return render_template("terms.html")


@app.route("/law")
def law():
    return render_template("law.html")


@app.route("/privacy")
def privacy():
    return render_template("privacy.html")


@app.route("/news")
def news():

    workbook = load_workbook("ニュース一覧.xlsx", data_only=True)
    sheet = workbook.active

    news_posts = []

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if not row[0]:
            continue

        news_posts.append({
            "id": row[0],
            "date": row[1],
            "category": row[2],
            "title": row[3],
            "slug": row[4],
            "summary": row[5],
            "image": row[6],
            "body": row[7],
            "seo_title": row[8],
            "seo_description": row[9],
            "related_product": row[10],
            "status": row[11]
        })

    news_posts = [
        post for post in news_posts
        if str(post["status"]).strip() == "公開"
    ]

    return render_template(
        "news.html",
        news_posts=news_posts
    )


@app.route("/news/<slug>")
def news_detail(slug):

    workbook = load_workbook("ニュース一覧.xlsx", data_only=True)
    sheet = workbook.active

    news_post = None

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if not row[0]:
            continue

        post = {
            "id": row[0],
            "date": row[1],
            "category": row[2],
            "title": row[3],
            "slug": row[4],
            "summary": row[5],
            "image": row[6],
            "body": row[7],
            "seo_title": row[8],
            "seo_description": row[9],
            "related_product": row[10],
            "status": row[11]
        }

        if (
            str(post["status"]).strip() == "公開"
            and str(post["slug"]).strip() == slug
        ):
            news_post = post
            break

    if news_post is None:
        return "ニュースが見つかりません", 404

    related_product_id = None

    if news_post["related_product"]:
        related_product_name = str(
            news_post["related_product"]
        ).strip()

        for product_id, product in products_from_excel.items():

            if product["name"].strip() == related_product_name:
                related_product_id = product_id
                break

    return render_template(
        "news_detail.html",
        news_post=news_post,
        related_product_id=related_product_id
    )
@app.route("/business")
def business():
    return render_template("business.html")

@app.route('/blog')
def blog():

    workbook = load_workbook("ブログ一覧.xlsx", data_only=True)
    sheet = workbook.active

    blog_posts = []

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if not row[0]:
            continue

        blog_posts.append({
            "id": row[0],
            "date": row[1],
            "category": row[2],
            "title": row[3],
            "slug": row[4],
            "summary": row[5],
            "image": row[6],
            "body": row[7],
            "seo_title": row[8],
            "seo_description": row[9],
            "related_product": row[10],
            "status": row[11]
        })

    blog_posts = [
        post for post in blog_posts
        if str(post["status"]).strip() == "公開"
    ]

    return render_template(
        "blog.html",
        blog_posts=blog_posts
    )
@app.route('/blog/<slug>')
def blog_detail(slug):

    workbook = load_workbook("ブログ一覧.xlsx", data_only=True)
    sheet = workbook.active

    blog_post = None

    for row in sheet.iter_rows(min_row=2, values_only=True):

        if not row[0]:
            continue

        post = {
            "id": row[0],
            "date": row[1],
            "category": row[2],
            "title": row[3],
            "slug": row[4],
            "summary": row[5],
            "image": row[6],
            "body": row[7],
            "seo_title": row[8],
            "seo_description": row[9],
            "related_product": row[10],
            "status": row[11]
        }

        if (
            str(post["status"]).strip() == "公開"
            and str(post["slug"]).strip() == slug
        ):
            blog_post = post
            break

    if blog_post is None:
        return "記事が見つかりません", 404

    return render_template(
        "blog_detail.html",
        blog_post=blog_post
    )
# =========================
# ログイン
# =========================

@app.route("/login")
def login():
    return render_template("login.html")


# =========================
# アプリ起動
# =========================

if __name__ == "__main__":
    app.run(debug=True, port=5001)